import pptx, json, sys

def parse_pptx(path):
    prs = pptx.Presentation(path)
    slides_data = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells]
                    texts.append(' | '.join(row_texts))
        slides_data.append({'slide': i+1, 'texts': texts})
    return slides_data

if __name__ == '__main__':
    path = sys.argv[1]
    data = parse_pptx(path)
    print(json.dumps(data, ensure_ascii=False, indent=2))