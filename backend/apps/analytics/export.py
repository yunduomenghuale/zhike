"""班级学习数据 Excel 导出（教师端学情分析页「导出 Excel」，2026-10-29）。

口径约定（与教师确认）：
- 教师视角全量导出：不受考试 score_released / 作业 correct_status 门控限制，
  未发布成绩照实导出（教师是评分主体，导出即教师工作用途）。
- 折算分：单次折算分 = 得分 ÷ 该次满分 × 100（保留 1 位小数）；
  「折算平均分」= 该生各次折算分的算术平均（每次任务等权，避免满分大的
  任务主导均分）；「原始平均分」= 各次原始得分的算术平均。
- 时间范围（start/end）作用于「行为发生时间」：练习 submitted_at、作业/考试/
  实验 submitted_at、问答 created_at、视频进度 updated_at；
  应交/应考/应做任务清单不受时间过滤（任务发布即存在）。
- 缺考/缺交行保留，成绩列留空，方便教师一眼看到谁未参加。

数据全部现查现算，不复用 build_class_stats：该函数无时间过滤参数，
且导出需要的是明细行而非聚合行。
"""
from datetime import datetime, time, timedelta

from django.db.models import Count, Q
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.classroom.models import ClassStudent
from apps.courses.models import Catalog, TeachingVideo, VideoWatchProgress
from apps.exams.models import Exam, ExamLog, ExamSubmission
from apps.homework.models import Homework, HomeworkSubmission
from apps.knowledge.models import QARecord
from apps.labs.models import Lab, LabSubmission
from apps.questions.models import AnswerRecord
from apps.users.models import User

from .services import _chapter_maps

# ---- 样式（扁平、打印友好；不依赖主题） ----
_HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
_TITLE_FONT = Font(size=13, bold=True, color="1F3864")
_HEADER_FONT = Font(bold=True, color="1F3864")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def parse_export_range(params):
    """解析 start/end 查询参数。返回 (start_dt, end_dt, error_message)。

    支持 YYYY-MM-DD 或 ISO 日期时间；纯日期 end 取当日末尾（含当天）。
    不传 = 全量（None, None）。
    """
    start_raw = (params.get("start") or "").strip()
    end_raw = (params.get("end") or "").strip()

    def _parse(raw, is_end):
        if not raw:
            return None, None
        # 先按纯日期解析：parse_datetime 对 "2026-10-31" 也会成功但返回零点，
        # 会让 is_end 的「含当天」分支失效；日期串必须走 combine(time.max)。
        d = parse_date(raw)
        if d is not None:
            dt = datetime.combine(d, time.max if is_end else time.min)
        else:
            dt = parse_datetime(raw)
            if dt is None:
                return None, "时间格式不正确，应为 YYYY-MM-DD 或 ISO 日期时间"
        if timezone.is_naive(dt):
            dt = timezone.make_aware(dt, timezone.get_current_timezone())
        return dt, None

    start, err = _parse(start_raw, False)
    if err:
        return None, None, err
    end, err = _parse(end_raw, True)
    if err:
        return None, None, err
    if start and end and start > end:
        return None, None, "开始时间不能晚于结束时间"
    return start, end, None


def _pct(score, total):
    """折算分：得分 ÷ 满分 × 100，保留 1 位小数；缺任一项返回 None。"""
    if score is None or not total:
        return None
    return round(float(score) / float(total) * 100, 1)


def _f1(value):
    """数值统一 1 位小数；None 原样返回（Excel 里显示空单元格）。"""
    return round(float(value), 1) if value is not None else None


def _fmt_dt(dt):
    return timezone.localtime(dt).strftime("%Y-%m-%d %H:%M") if dt else ""


def _fmt_minutes(seconds):
    if seconds is None:
        return None
    return round(float(seconds) / 60, 1)


class _SheetWriter:
    """统一的 Sheet 写法：第 1 行标题、第 2 行表头、数据从第 3 行起，冻结表头。"""

    def __init__(self, wb, title, sheet_name, headers):
        self.ws = wb.create_sheet(sheet_name)
        self.headers = headers
        n = len(headers)
        self.ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        c = self.ws.cell(row=1, column=1, value=title)
        c.font = _TITLE_FONT
        c.alignment = Alignment(vertical="center")
        self.ws.row_dimensions[1].height = 24
        for i, h in enumerate(headers, 1):
            cell = self.ws.cell(row=2, column=i, value=h)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _CENTER
            cell.border = _BORDER
        self.ws.freeze_panes = "A3"
        self._row = 3

    def append(self, values):
        if len(values) != len(self.headers):
            raise ValueError(
                f"Sheet「{self.ws.title}」数据列数 {len(values)} 与表头 {len(self.headers)} 不一致"
            )
        for i, v in enumerate(values, 1):
            cell = self.ws.cell(row=self._row, column=i, value=v)
            cell.border = _BORDER
            if isinstance(v, (int, float)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
        self._row += 1

    def autosize(self, sample_rows=200):
        """按表头 + 前若干行内容估算列宽（中文字符按 2 计）。"""
        for idx, h in enumerate(self.headers, 1):
            width = self._text_width(h) + 4
            for row in self.ws.iter_rows(min_row=3, max_row=min(self._row - 1, 2 + sample_rows),
                                         min_col=idx, max_col=idx):
                for cell in row:
                    if cell.value is not None:
                        width = max(width, self._text_width(str(cell.value)) + 2)
            self.ws.column_dimensions[get_column_letter(idx)].width = min(width, 42)

    @staticmethod
    def _text_width(text):
        return sum(2 if ord(ch) > 127 else 1 for ch in text)


def _range_filter(field_name, start, end):
    """构造 ORM 时间范围过滤 Q 对象。"""
    q = Q()
    if start:
        q &= Q(**{f"{field_name}__gte": start})
    if end:
        q &= Q(**{f"{field_name}__lte": end})
    return q


# 可勾选的导出内容（key → Sheet 名）。overview 恒导出（其余 Sheet 是它的明细下钻）。
EXPORT_SECTIONS = [
    ("overview", "成绩总览"),
    ("homework", "作业明细"),
    ("exam", "考试明细"),
    ("lab", "实验明细"),
    ("practice", "练习与错题"),
    ("qa", "问答与视频"),
    ("chapter", "章节矩阵"),
]
SECTION_KEYS = [key for key, _ in EXPORT_SECTIONS]
ALWAYS_SECTIONS = {"overview"}


def parse_sections(raw):
    """解析 sheets 参数（逗号分隔）。返回 (sections, error)。

    未传 = 全部；非法 key 报错；overview 恒包含（教师总要有一张汇总表）。
    """
    raw = (raw or "").strip()
    if not raw:
        return set(SECTION_KEYS), None
    keys = {k.strip() for k in raw.split(",") if k.strip()}
    unknown = keys - set(SECTION_KEYS)
    if unknown:
        return None, f"未知的导出内容：{'、'.join(sorted(unknown))}"
    return keys | ALWAYS_SECTIONS, None


def build_class_export_workbook(classroom, course, start=None, end=None, sections=None) -> Workbook:
    """生成班级 × 课程的学习数据工作簿（按 sections 裁剪 Sheet，口径见模块 docstring）。"""
    sections = set(sections or SECTION_KEYS) | ALWAYS_SECTIONS
    wb = Workbook()
    wb.remove(wb.active)  # 去掉默认 Sheet

    now = timezone.localtime()
    range_text = "全部时间" if not (start or end) else (
        f"{_fmt_dt(start) or '起始'} ~ {_fmt_dt(end) or '至今'}"
    )
    base_title = f"{classroom.name} · 《{course.name}》 学习数据导出"
    subtitle = f"时间范围：{range_text}　导出时间：{now.strftime('%Y-%m-%d %H:%M:%S')}"

    # 导出名单与教师端学情统计同口径（方案 §5.4）：占用关系且在读。
    # 毕业/休学学生不占导出行——他们的未完成任务会把班级完成率永久拉低，
    # 而导出的核心正是完成率口径；历史成绩由学生本人端保留可查。
    students = list(
        ClassStudent.objects.filter(
            classroom=classroom,
            learn_status=ClassStudent.LearnStatus.ACTIVE,
            student__edu_status=User.EduStatus.ENROLLED,
        )
        .select_related("student")
        .order_by("student__username")
    )
    stu_meta = [
        (cs.student_id, cs.student.username, cs.student.real_name or cs.student.username)
        for cs in students
    ]
    stu_ids = [m[0] for m in stu_meta]

    # ---- 任务清单（不受时间过滤：发布即应做/应考） ----
    hw_list = list(Homework.objects.filter(
        classroom=classroom, course=course, status=Homework.Status.PUBLISHED
    ).order_by("id"))
    exam_list = list(Exam.objects.filter(
        classroom=classroom, course=course,
        status__in=[Exam.Status.PUBLISHED, Exam.Status.FINISHED],
    ).order_by("id"))
    lab_list = list(Lab.objects.filter(
        course=course, status=Lab.Status.PUBLISHED, schedules__classroom=classroom
    ).distinct().order_by("order", "id"))

    # ---- 明细数据（一次查出，内存中组装矩阵） ----
    hw_subs = {
        (s.homework_id, s.student_id): s
        for s in HomeworkSubmission.objects.filter(
            student_id__in=stu_ids, homework__classroom=classroom, homework__course=course
        ).filter(_range_filter("submitted_at", start, end))
    }
    exam_subs = {
        (s.exam_id, s.student_id): s
        for s in ExamSubmission.objects.filter(
            student_id__in=stu_ids, exam__classroom=classroom, exam__course=course
        ).filter(_range_filter("submitted_at", start, end))
    }
    lab_subs = {
        (s.lab_id, s.student_id): s
        for s in LabSubmission.objects.filter(
            student_id__in=stu_ids, lab_id__in=[lab.id for lab in lab_list],
            status=LabSubmission.Status.SUBMITTED,
        ).filter(_range_filter("submitted_at", start, end))
    }
    exam_log_counts = {
        (row["exam_id"], row["student_id"]): row["n"]
        for row in ExamLog.objects.filter(exam__in=exam_list, student_id__in=stu_ids)
        .values("exam_id", "student_id").annotate(n=Count("id"))
    }
    practice_qs = AnswerRecord.objects.filter(
        student_id__in=stu_ids, scene=AnswerRecord.Scene.PRACTICE, question__course_id=course.id
    ).filter(_range_filter("submitted_at", start, end))
    qa_qs = QARecord.objects.filter(student_id__in=stu_ids, course=course).filter(
        _range_filter("created_at", start, end)
    )
    video_qs = VideoWatchProgress.objects.filter(
        student_id__in=stu_ids, video__course_id=course.id
    ).filter(_range_filter("updated_at", start, end))

    # 逐生汇总（总览/练习/问答视频三表共用）
    agg = {sid: {
        "practice_total": 0, "practice_correct": 0, "wrong_qids": set(),
        "qa_count": 0, "qa_disliked": 0, "qa_by_chapter": {},
        "watch_seconds": 0, "video_done": 0, "video_progress": 0, "video_not_started": 0,
        "hw_scores": [], "hw_pcts": [],
        "exam_scores": [], "exam_pcts": [], "exam_minutes": [],
        "lab_scores": [], "lab_pcts": [],
    } for sid in stu_ids}

    for row in practice_qs.values("student_id", "is_correct", "question_id"):
        a = agg[row["student_id"]]
        a["practice_total"] += 1
        if row["is_correct"]:
            a["practice_correct"] += 1
        else:
            a["wrong_qids"].add(row["question_id"])

    _parent_of, title_of, root_id = _chapter_maps(course.id)

    def root_title(cid):
        rid = root_id(cid)
        return title_of.get(rid, "未分类章节") if rid else "课程级"

    for row in qa_qs.values("student_id", "catalog_id", "feedback"):
        a = agg[row["student_id"]]
        a["qa_count"] += 1
        if row["feedback"] == -1:
            a["qa_disliked"] += 1
        chapter = root_title(row["catalog_id"]) if row["catalog_id"] else "课程级"
        a["qa_by_chapter"][chapter] = a["qa_by_chapter"].get(chapter, 0) + 1

    for row in video_qs.values("student_id", "watch_seconds", "status"):
        a = agg[row["student_id"]]
        a["watch_seconds"] += row["watch_seconds"] or 0
        if row["status"] == VideoWatchProgress.Status.COMPLETED:
            a["video_done"] += 1
        elif row["status"] == VideoWatchProgress.Status.IN_PROGRESS:
            a["video_progress"] += 1
        else:
            a["video_not_started"] += 1

    for (hw_id, sid), sub in hw_subs.items():
        hw = next(h for h in hw_list if h.id == hw_id)
        if sub.score is None:
            continue
        a = agg[sid]
        a["hw_scores"].append(float(sub.score))
        pct = _pct(sub.score, hw.total_score)
        if pct is not None:
            a["hw_pcts"].append(pct)
    for (exam_id, sid), sub in exam_subs.items():
        if sub.total_score is None:
            continue
        exam = next(e for e in exam_list if e.id == exam_id)
        a = agg[sid]
        a["exam_scores"].append(float(sub.total_score))
        pct = _pct(sub.total_score, exam.total_score)
        if pct is not None:
            a["exam_pcts"].append(pct)
        if sub.started_at and sub.submitted_at:
            a["exam_minutes"].append((sub.submitted_at - sub.started_at).total_seconds() / 60)
    for (lab_id, sid), sub in lab_subs.items():
        if sub.total_score is None:
            continue
        lab = next(l for l in lab_list if l.id == lab_id)
        a = agg[sid]
        a["lab_scores"].append(float(sub.total_score))
        pct = _pct(sub.total_score, lab.total_score)
        if pct is not None:
            a["lab_pcts"].append(pct)

    def _avg(seq):
        return round(sum(seq) / len(seq), 1) if seq else None

    # ============ Sheet 1 成绩总览 ============
    ws1 = _SheetWriter(wb, f"{base_title}（{subtitle}）", "成绩总览", [
        "学号", "姓名",
        "练习答题量", "答对数", "正确率(%)",
        "作业应交", "已交", "提交率(%)", "作业平均分(原始)", "作业折算平均分",
        "考试应考", "已考", "考试平均分(原始)", "考试折算平均分", "考试平均用时(分钟)",
        "实验应做", "已做", "实验平均分(原始)", "实验折算平均分",
        "观看时长(分钟)", "视频完成", "视频总数", "提问数", "被点无用",
    ])
    # 视频总数是分母（课程全部教学视频），不受时间过滤；观看/完成数受过滤
    total_videos = TeachingVideo.objects.filter(course=course).count()
    hw_total, exam_total, lab_total = len(hw_list), len(exam_list), len(lab_list)
    for sid, username, name in stu_meta:
        a = agg[sid]
        p_total, p_correct = a["practice_total"], a["practice_correct"]
        hw_sub_n = len({hid for (hid, s) in hw_subs if s == sid})
        exam_taken_n = len({eid for (eid, s) in exam_subs if s == sid})
        lab_done_n = len({lid for (lid, s) in lab_subs if s == sid})
        ws1.append([
            username, name,
            p_total, p_correct,
            round(p_correct / p_total * 100, 1) if p_total else None,
            hw_total, hw_sub_n,
            round(hw_sub_n / hw_total * 100, 1) if hw_total else None,
            _avg(a["hw_scores"]), _avg(a["hw_pcts"]),
            exam_total, exam_taken_n,
            _avg(a["exam_scores"]), _avg(a["exam_pcts"]), _avg(a["exam_minutes"]),
            lab_total, lab_done_n,
            _avg(a["lab_scores"]), _avg(a["lab_pcts"]),
            _fmt_minutes(a["watch_seconds"]), a["video_done"], total_videos,
            a["qa_count"], a["qa_disliked"],
        ])
    ws1.autosize()

    # ============ Sheet 2 作业明细（作业 × 学生 全矩阵，缺交留空） ============
    if "homework" in sections:
        ws2 = _SheetWriter(wb, base_title, "作业明细", [
            "学号", "姓名", "作业名称", "满分", "得分", "折算分",
            "是否迟交", "提交时间", "批改状态", "提交情况",
        ])
        for hw in hw_list:
            for sid, username, name in stu_meta:
                sub = hw_subs.get((hw.id, sid))
                ws2.append([
                    username, name, hw.title,
                    float(hw.total_score),
                    _f1(sub.score) if sub else None,
                    _pct(sub.score, hw.total_score) if sub else None,
                    "是" if sub and sub.is_late else ("否" if sub else None),
                    _fmt_dt(sub.submitted_at) if sub else None,
                    sub.get_correct_status_display() if sub else None,
                    "已交" if sub else "缺交",
                ])
        ws2.autosize()

    # ============ Sheet 3 考试明细（考试 × 学生 全矩阵） ============
    if "exam" in sections:
        ws3 = _SheetWriter(wb, base_title, "考试明细", [
            "学号", "姓名", "考试名称", "满分", "总分", "折算分",
            "客观题得分", "主观题得分", "用时(分钟)", "状态", "异常操作次数", "成绩已发布",
        ])
        for exam in exam_list:
            for sid, username, name in stu_meta:
                sub = exam_subs.get((exam.id, sid))
                subjective = None
                if sub and isinstance(sub.subjective_scores, dict) and sub.subjective_scores:
                    subjective = round(sum(float(v) for v in sub.subjective_scores.values() if v is not None), 1)
                minutes = None
                if sub and sub.started_at and sub.submitted_at:
                    minutes = round((sub.submitted_at - sub.started_at).total_seconds() / 60, 1)
                ws3.append([
                    username, name, exam.name,
                    float(exam.total_score),
                    _f1(sub.total_score) if sub else None,
                    _pct(sub.total_score, exam.total_score) if sub else None,
                    _f1(sub.objective_score) if sub else None,
                    subjective,
                    minutes,
                    sub.get_status_display() if sub else "缺考",
                    exam_log_counts.get((exam.id, sid), 0) if sub else None,
                    ("是" if sub.score_released else "否") if sub else None,
                ])
        ws3.autosize()

    # ============ Sheet 4 实验明细（实验 × 学生 全矩阵） ============
    if "lab" in sections:
        ws4 = _SheetWriter(wb, base_title, "实验明细", [
            "学号", "姓名", "实验名称", "满分", "总分", "折算分",
            "基础分", "操作准确性", "完成效率", "完成度",
            "通过步骤", "总步骤", "用时(分钟)", "第几次作答", "教师复核", "状态",
        ])
        for lab in lab_list:
            for sid, username, name in stu_meta:
                sub = lab_subs.get((lab.id, sid))
                ws4.append([
                    username, name, lab.title,
                    float(lab.total_score),
                    _f1(sub.total_score) if sub else None,
                    _pct(sub.total_score, lab.total_score) if sub else None,
                    _f1(sub.base_score) if sub else None,
                    _f1(sub.accuracy_score) if sub else None,
                    _f1(sub.efficiency_score) if sub else None,
                    _f1(sub.completion_score) if sub else None,
                    sub.passed_steps if sub else None,
                    sub.total_steps if sub else None,
                    _fmt_minutes(sub.elapsed_seconds) if sub else None,
                    sub.attempt if sub else None,
                    ("是" if sub.reviewed else "否") if sub else None,
                    "已提交" if sub else "未提交",
                ])
        ws4.autosize()

    # ============ Sheet 5 练习与错题（每生一行） ============
    if "practice" in sections:
        ws5 = _SheetWriter(wb, base_title, "练习与错题", [
            "学号", "姓名", "答题量", "答对", "答错", "正确率(%)", "错题数(按题目去重)",
        ])
        for sid, username, name in stu_meta:
            a = agg[sid]
            p_total, p_correct = a["practice_total"], a["practice_correct"]
            ws5.append([
                username, name, p_total, p_correct, p_total - p_correct,
                round(p_correct / p_total * 100, 1) if p_total else None,
                len(a["wrong_qids"]),
            ])
        ws5.autosize()

    # ============ Sheet 6 问答与视频（每生一行） ============
    if "qa" in sections:
        ws6 = _SheetWriter(wb, base_title, "问答与视频", [
            "学号", "姓名", "提问总数", "被点无用", "提问章节分布",
            "观看时长(分钟)", "已完成章节", "学习中章节", "未开始章节",
        ])
        for sid, username, name in stu_meta:
            a = agg[sid]
            qa_dist = "；".join(f"{ch}: {n}" for ch, n in sorted(
                a["qa_by_chapter"].items(), key=lambda kv: -kv[1]
            ))
            ws6.append([
                username, name, a["qa_count"], a["qa_disliked"], qa_dist or None,
                _fmt_minutes(a["watch_seconds"]),
                a["video_done"], a["video_progress"], a["video_not_started"],
            ])
        ws6.autosize()

    # ============ Sheet 7 章节矩阵（学生 × 顶层章节 正确率%） ============
    if "chapter" in sections:
        top_chapters = [
            c for c in Catalog.objects.filter(course=course, parent__isnull=True, is_published=True)
            .order_by("order", "id")
        ]
        chapter_acc = {sid: {c.id: [0, 0] for c in top_chapters} for sid in stu_ids}  # [correct, total]
        for row in practice_qs.values("student_id", "question__catalog_id", "is_correct"):
            rid = root_id(row["question__catalog_id"])
            if rid in chapter_acc[row["student_id"]]:
                cell = chapter_acc[row["student_id"]][rid]
                cell[1] += 1
                if row["is_correct"]:
                    cell[0] += 1
        ws7 = _SheetWriter(
            wb, base_title, "章节矩阵",
            ["学号", "姓名"] + [c.title for c in top_chapters],
        )
        for sid, username, name in stu_meta:
            cells = []
            for c in top_chapters:
                correct, total = chapter_acc[sid][c.id]
                cells.append(round(correct / total * 100, 1) if total else None)
            ws7.append([username, name] + cells)
        ws7.autosize()

    return wb
