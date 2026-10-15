import json
import os
import time
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase, override_settings
from pptx import Presentation
from pptx.util import Inches

from apps.ai.providers.mock import MockProvider
from apps.ai.services import (
    _dedupe_page_body,
    _local_script_for_page,
    _polish_script_text,
    generate_script_pages_batched,
    ingest_ppt_pages,
    synthesize_audio_for_video,
)
from apps.courses.ppt_parser import _read_presentation
from apps.knowledge.models import KnowledgeChunk

User = get_user_model()


class TeachingScriptCleanupTests(SimpleTestCase):
    def test_page_body_removes_numbered_and_exact_duplicate_lines(self):
        body = "1. Application\n2. Application\nApplet\nApplet"

        self.assertEqual(_dedupe_page_body(body), "1. Application\nApplet")

    def test_polish_removes_repeated_generated_clause(self):
        sentence = "Application 是普通 Java 应用程序，通常从 main 方法开始运行"
        script = f"{sentence}；{sentence}；Applet 是早期嵌入网页运行的 Java 程序。"

        polished = _polish_script_text(script)

        self.assertEqual(polished.count(sentence), 1)
        self.assertIn("Applet 是早期嵌入网页运行的 Java 程序", polished)

    def test_local_script_deduplicates_equivalent_explanations_and_has_no_fixed_filler(self):
        result = _local_script_for_page(
            {
                "page": 6,
                "title": "Application 和 Applet",
                "body": "Application\nApplication 程序\nApplet",
            }
        )

        script = result["script"]
        self.assertEqual(script.count("通常从 main 方法开始运行"), 1)
        self.assertNotIn("学习时先把这些要点按顺序串起来", script)

    def test_ppt_parser_ignores_duplicate_text_boxes(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "duplicate-text.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Application 和 Applet"
            slide.placeholders[1].text = "Application\nApplet"
            duplicate = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1))
            duplicate.text_frame.text = "Application\nApplet"
            presentation.save(path)

            pages = _read_presentation(str(path))

        self.assertEqual(pages[0]["body"], "Application\nApplet")


class _FakeGenStatus:
    SCRIPT_READY = "script_ready"
    AUDIO_READY = "audio_ready"


class _FakeVideo:
    """教学视频替身：只保留生成函数用到的属性，避开数据库。"""

    GenStatus = _FakeGenStatus

    def __init__(self, scripts):
        self.scripts = scripts
        self.gen_status = None
        self.is_published = False
        self.published_at = None
        self.save_calls = []

    def save(self, update_fields=None):
        self.save_calls.append(update_fields)


class _FakeChat:
    """假 chat provider：记录调用参数，按请求内嵌的页面列表返回 JSON。"""

    name = "fake-chat"

    def __init__(self, delay=0.4):
        self.delay = delay
        self.calls = []

    def chat(self, messages, **kwargs):
        self.calls.append(kwargs)
        time.sleep(self.delay)
        payload = messages[1]["content"].split("PPT 页面：\n", 1)[1]
        pages = json.loads(payload)
        return json.dumps(
            [{"page": p["page"], "script": f"第{p['page']}页的课堂讲解稿内容。"} for p in pages],
            ensure_ascii=False,
        )


class _FakeTTS:
    """假 tts provider：可配置延迟与失败标记。"""

    name = "fake-tts"

    def __init__(self, delay=0.3, fail_marker=None):
        self.delay = delay
        self.fail_marker = fail_marker

    def tts(self, text, voice="Cherry"):
        time.sleep(self.delay)
        if self.fail_marker and self.fail_marker in text:
            raise RuntimeError("synthetic tts failure")
        return f"/media/tts/fake-{abs(hash(text)) % 10**8}.wav"


class GenerationConcurrencyTests(SimpleTestCase):
    """讲稿/配音并发改造的行为与性能回归测试。"""

    # ---------- 讲稿：单请求双批并发 ----------

    def test_script_two_chunks_run_concurrently_and_merge_by_page(self):
        pages = [{"page": i, "title": f"标题{i}", "body": f"第{i}页正文内容。"} for i in range(1, 13)]
        video = _FakeVideo([])
        chat = _FakeChat(delay=0.4)
        with patch("apps.ai.services.get_provider", return_value=chat):
            started = time.monotonic()
            result = generate_script_pages_batched(video, pages, batch_size=6)
            elapsed = time.monotonic() - started

        # 12 页 = 2 个 6 页块并发：耗时应接近单块（0.4s），而非串行的 0.8s
        self.assertLess(elapsed, 0.72)
        self.assertEqual(len(chat.calls), 2)
        self.assertEqual([call.get("timeout") for call in chat.calls], [45, 45])
        self.assertEqual([call.get("retries") for call in chat.calls], [2, 2])
        # 全部 12 页按页码归位、无遗漏无串台
        self.assertEqual(len(video.scripts), 12)
        self.assertTrue(all(item.get("script") for item in video.scripts))
        self.assertEqual([item["page"] for item in video.scripts], list(range(1, 13)))
        self.assertTrue(result["done"])
        # 初次落库 + 全部完成后的发布状态落库，共两次保存
        self.assertEqual(len(video.save_calls), 2)
        self.assertEqual(video.gen_status, _FakeGenStatus.SCRIPT_READY)
        self.assertTrue(video.is_published)

    def test_script_ai_failure_fills_local_fallback(self):
        pages = [{"page": i, "title": f"标题{i}", "body": f"第{i}页正文。"} for i in range(1, 4)]

        class _Boom:
            name = "fake-chat"

            def chat(self, messages, **kwargs):
                raise RuntimeError("provider down")

        video = _FakeVideo([])
        with patch("apps.ai.services.get_provider", return_value=_Boom()):
            result = generate_script_pages_batched(video, pages, batch_size=6)

        # AI 整批失败：本地兜底文本顶上，页面不留空
        self.assertTrue(result["done"])
        self.assertTrue(all(str(item.get("script") or "").strip() for item in video.scripts))

    def test_script_skips_pages_already_generated(self):
        pages = [{"page": i, "title": f"标题{i}", "body": f"第{i}页正文。"} for i in range(1, 5)]
        video = _FakeVideo([
            {"page": 1, "script": "已有讲解稿", "audio_url": "/media/tts/a.wav"},
            {"page": 2, "script": "已有讲解稿2"},
            {"page": 3, "script": ""},
            {"page": 4, "script": ""},
        ])
        chat = _FakeChat(delay=0.01)
        with patch("apps.ai.services.get_provider", return_value=chat):
            result = generate_script_pages_batched(video, pages, batch_size=6)

        # 已有讲解稿的页不被重算、已有 audio_url 不被覆盖
        self.assertEqual(video.scripts[0]["script"], "已有讲解稿")
        self.assertEqual(video.scripts[0]["audio_url"], "/media/tts/a.wav")
        self.assertEqual(result["script_pages"], 4)
        self.assertTrue(result["done"])

    # ---------- 配音：批内并发 ----------

    def test_synthesize_runs_concurrently_and_saves_once(self):
        scripts = [{"page": i, "script": f"第{i}页的讲解稿文本。"} for i in range(1, 5)]
        video = _FakeVideo(scripts)
        tts = _FakeTTS(delay=0.3)
        with patch("apps.ai.services.get_provider", return_value=tts):
            started = time.monotonic()
            ok = synthesize_audio_for_video(video, limit=3)
            elapsed = time.monotonic() - started

        self.assertEqual(ok, 3)
        # 3 页并发：耗时接近单页（0.3s），而非串行 0.9s
        self.assertLess(elapsed, 0.7)
        self.assertTrue(all(scripts[i].get("audio_url") for i in range(3)))
        self.assertFalse(scripts[3].get("audio_url"))
        # 仅批末一次写库（原为每页一次）
        self.assertEqual(len(video.save_calls), 1)
        # 未完成全部配音：gen_status 仍为 SCRIPT_READY、不发布
        self.assertEqual(video.gen_status, _FakeGenStatus.SCRIPT_READY)
        self.assertFalse(video.is_published)

    def test_synthesize_partial_failure_recorded_and_publish_when_done(self):
        scripts = [
            {"page": 1, "script": "第一页正常讲解稿。"},
            {"page": 2, "script": "第二页会失败的讲解稿。"},
        ]
        video = _FakeVideo(scripts)
        tts = _FakeTTS(delay=0.01, fail_marker="会失败")
        with patch("apps.ai.services.get_provider", return_value=tts):
            ok = synthesize_audio_for_video(video, limit=2)

        self.assertEqual(ok, 1)
        self.assertTrue(scripts[0].get("audio_url"))
        self.assertIn("synthetic tts failure", scripts[2 - 1]["audio_error"])
        self.assertEqual(video.gen_status, _FakeGenStatus.SCRIPT_READY)

        # 修复后补跑：失败页增量补齐，全部完成即自动发布
        tts2 = _FakeTTS(delay=0.01)
        with patch("apps.ai.services.get_provider", return_value=tts2):
            ok2 = synthesize_audio_for_video(video, limit=2)
        self.assertEqual(ok2, 1)
        self.assertEqual(video.gen_status, _FakeGenStatus.AUDIO_READY)
        self.assertTrue(video.is_published)


class LocalModelTimeoutTests(SimpleTestCase):
    """本地模型场景：讲稿/TTS 超时与默认发音人均来自 settings（环境变量可调）。"""

    def test_script_batch_timeout_from_settings(self):
        pages = [{"page": 1, "title": "标题", "body": "正文。"}]
        video = _FakeVideo([])
        chat = _FakeChat(delay=0.01)
        with patch("apps.ai.services.get_provider", return_value=chat), override_settings(AI_SCRIPT_TIMEOUT=180):
            generate_script_pages_batched(video, pages, batch_size=6)

        # 本地模型推理慢：超时按配置调大后透传到 provider.chat
        self.assertEqual([call.get("timeout") for call in chat.calls], [180])

    def test_synthesize_default_voice_from_settings(self):
        scripts = [{"page": 1, "script": "第一页讲解稿。"}]
        video = _FakeVideo(scripts)

        class _VoiceRecorder(_FakeTTS):
            def __init__(self):
                super().__init__(delay=0.01)
                self.voice = None

            def tts(self, text, voice="Cherry"):
                self.voice = voice
                return "/media/tts/x.wav"

        tts = _VoiceRecorder()
        with patch("apps.ai.services.get_provider", return_value=tts), override_settings(AI_TTS_VOICE="local-speaker"):
            synthesize_audio_for_video(video, limit=1)

        self.assertEqual(tts.voice, "local-speaker")

    def test_synthesize_explicit_voice_wins_over_settings(self):
        scripts = [{"page": 1, "script": "第一页讲解稿。"}]
        video = _FakeVideo(scripts)

        class _VoiceRecorder(_FakeTTS):
            def __init__(self):
                super().__init__(delay=0.01)
                self.voice = None

            def tts(self, text, voice="Cherry"):
                self.voice = voice
                return "/media/tts/x.wav"

        tts = _VoiceRecorder()
        with patch("apps.ai.services.get_provider", return_value=tts), override_settings(AI_TTS_VOICE="local-speaker"):
            synthesize_audio_for_video(video, voice="Cherry", limit=1)

        self.assertEqual(tts.voice, "Cherry")

    def test_script_timeout_prefers_provider_config(self):
        """管理端给对话能力配了超时：优先于 settings.AI_SCRIPT_TIMEOUT。"""
        pages = [{"page": 1, "title": "标题", "body": "正文。"}]
        video = _FakeVideo([])
        chat = _FakeChat(delay=0.01)
        chat.timeout = 180  # 管理端配置经 build_provider 物化到 provider 上
        with patch("apps.ai.services.get_provider", return_value=chat), override_settings(AI_SCRIPT_TIMEOUT=45):
            generate_script_pages_batched(video, pages, batch_size=6)

        self.assertEqual([call.get("timeout") for call in chat.calls], [180])

    def test_synthesize_voice_prefers_provider_config(self):
        """管理端给语音能力配了发音人：优先于 settings.AI_TTS_VOICE。"""
        scripts = [{"page": 1, "script": "第一页讲解稿。"}]
        video = _FakeVideo(scripts)

        class _VoiceRecorder(_FakeTTS):
            def __init__(self):
                super().__init__(delay=0.01)
                self.voice = None

            def tts(self, text, voice="Cherry"):
                self.voice = voice
                return "/media/tts/x.wav"

        tts = _VoiceRecorder()
        tts.voice = "speaker-01"  # 管理端配置
        with patch("apps.ai.services.get_provider", return_value=tts), override_settings(AI_TTS_VOICE="local-speaker"):
            synthesize_audio_for_video(video, limit=1)

        self.assertEqual(tts.voice, "speaker-01")


class _FakeTTSResponse:
    """假 HTTP 响应：供 OpenAICompatProvider TTS 分支测试使用。"""

    def __init__(self, *, status_code=200, content=b"", headers=None, json_body=None):
        self.status_code = status_code
        self.content = content
        self.headers = headers or {}
        self._json_body = json_body

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")

    def json(self):
        if self._json_body is None:
            raise ValueError("响应体不是 JSON")
        return self._json_body


def _local_tts_provider(base_url="http://127.0.0.1:8000/v1", tts_model="tts-local"):
    from apps.ai.providers.openai_compat import OpenAICompatProvider

    return OpenAICompatProvider(
        {
            "api_key": "local-key",
            "base_url": base_url,
            "chat_model": "local-chat",
            "tts_model": tts_model,
        }
    )


class OpenAICompatibleTTSTests(SimpleTestCase):
    """OpenAI 兼容 TTS 分支（本地 localai/AllTalk 等）行为测试。"""

    def test_posts_openai_payload_and_saves_audio(self):
        provider = _local_tts_provider()
        seen = {}

        def fake_post(url, headers=None, json=None, timeout=None):
            seen.update(url=url, json=json, timeout=timeout)
            return _FakeTTSResponse(content=b"ID3fake-mp3", headers={"Content-Type": "audio/mpeg"})

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                url = provider.tts("第一页讲解稿。", voice="speaker-a")
                # 音频二进制直接落盘，返回 media 可访问路径
                self.assertTrue(url.startswith("/media/tts/"))
                with open(os.path.join(tmp, url[len("/media/"):]), "rb") as f:
                    self.assertEqual(f.read(), b"ID3fake-mp3")

        # 打的是 OpenAI 协议 /audio/speech，payload 含 model/input/voice/response_format
        self.assertEqual(seen["url"], "http://127.0.0.1:8000/v1/audio/speech")
        self.assertEqual(seen["json"]["model"], "tts-local")
        self.assertEqual(seen["json"]["input"], "第一页讲解稿。")
        self.assertEqual(seen["json"]["voice"], "speaker-a")
        self.assertEqual(seen["json"]["response_format"], "mp3")

    def test_extension_follows_content_type(self):
        provider = _local_tts_provider()

        def fake_post(url, headers=None, json=None, timeout=None):
            return _FakeTTSResponse(content=b"RIFFfake-wav", headers={"Content-Type": "audio/x-wav"})

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                url = provider.tts("文本")

        self.assertTrue(url.endswith(".wav"))

    def test_4xx_raises_with_endpoint_hint_and_no_retry(self):
        provider = _local_tts_provider()
        calls = []

        def fake_post(url, headers=None, json=None, timeout=None):
            calls.append(url)
            return _FakeTTSResponse(
                status_code=404,
                content=b'{"error":{"message":"model not found","code":"model_not_found"}}',
                headers={"Content-Type": "application/json"},
                json_body={"error": {"message": "model not found", "code": "model_not_found"}},
            )

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                with self.assertRaises(RuntimeError) as ctx:
                    provider.tts("文本")

        # 错误信息带真实原因 + 端点（本地服务缺 /audio/speech 时能直接看出打到哪里）
        self.assertIn("model not found", str(ctx.exception))
        self.assertIn("/audio/speech", str(ctx.exception))
        self.assertEqual(len(calls), 1)  # 4xx 不重试
        self.assertIn("model not found", provider.last_tts_error)

    def test_connection_error_retries_three_times_then_raises(self):
        provider = _local_tts_provider()
        calls = []

        def fake_post(url, headers=None, json=None, timeout=None):
            calls.append(url)
            raise ConnectionError("local tts down")

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                with self.assertRaises(ConnectionError):
                    provider.tts("文本")

        self.assertEqual(len(calls), 3)
        self.assertIn("local tts down", provider.last_tts_error)

    def test_empty_audio_body_raises(self):
        provider = _local_tts_provider()

        def fake_post(url, headers=None, json=None, timeout=None):
            return _FakeTTSResponse(content=b"", headers={"Content-Type": "audio/mpeg"})

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                with self.assertRaises(RuntimeError) as ctx:
                    provider.tts("文本")

        self.assertIn("空音频", str(ctx.exception))

    def test_timeout_comes_from_settings(self):
        provider = _local_tts_provider()
        seen = {}

        def fake_post(url, headers=None, json=None, timeout=None):
            seen["timeout"] = timeout
            return _FakeTTSResponse(content=b"ID3", headers={"Content-Type": "audio/mpeg"})

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/", AI_TTS_TIMEOUT=321):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                provider.tts("文本")

        self.assertEqual(seen["timeout"], 321)

    def test_provider_timeout_and_voice_take_precedence(self):
        """管理端配置的 timeout/voice 优先于 settings 默认值（无需改 .env）。"""
        provider = _local_tts_provider()
        provider.timeout = 222
        provider.voice = "speaker-01"
        seen = {}

        def fake_post(url, headers=None, json=None, timeout=None):
            seen.update(json=json, timeout=timeout)
            return _FakeTTSResponse(content=b"ID3", headers={"Content-Type": "audio/mpeg"})

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                provider.tts("文本")

        self.assertEqual(seen["timeout"], 222)
        self.assertEqual(seen["json"]["voice"], "speaker-01")

    def test_dashscope_branch_uses_provider_voice(self):
        """DashScope 分支同样遵循管理端发音人配置。"""
        provider = _local_tts_provider(base_url="https://dashscope.aliyuncs.com/compatible-mode/v1")
        provider.voice = "Cherry"
        seen = {}
        native_body = {"output": {"audio": {"url": "https://example.com/a.wav"}}}
        native_content = json.dumps(native_body).encode()

        def fake_post(url, headers=None, json=None, timeout=None):
            seen["json"] = json
            return _FakeTTSResponse(
                content=native_content,
                headers={"Content-Type": "application/json"},
                json_body=native_body,
            )

        def fake_get(url, timeout=None):
            return _FakeTTSResponse(content=b"RIFFdata", headers={"Content-Type": "audio/wav"})

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                with patch("apps.ai.providers.openai_compat.requests.get", side_effect=fake_get):
                    provider.tts("你好")

        self.assertEqual(seen["json"]["input"]["voice"], "Cherry")

    def test_dashscope_base_url_still_uses_native_endpoint(self):
        """回归保护：DashScope 端点仍走原生协议（返回 URL 后下载落盘）。"""
        provider = _local_tts_provider(base_url="https://dashscope.aliyuncs.com/compatible-mode/v1")
        seen = {}
        native_body = {"output": {"audio": {"url": "https://example.com/a.wav"}}}
        native_content = json.dumps(native_body).encode()

        def fake_post(url, headers=None, json=None, timeout=None):
            seen["url"] = url
            seen["json"] = json
            return _FakeTTSResponse(
                content=native_content,
                headers={"Content-Type": "application/json"},
                json_body=native_body,
            )

        def fake_get(url, timeout=None):
            return _FakeTTSResponse(content=b"RIFFdata", headers={"Content-Type": "audio/wav"})

        with TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp, MEDIA_URL="/media/"):
            with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
                with patch("apps.ai.providers.openai_compat.requests.get", side_effect=fake_get):
                    url = provider.tts("你好", voice="Cherry")

        self.assertIn("/api/v1/services/aigc/multimodal-generation/generation", seen["url"])
        self.assertEqual(seen["json"]["input"]["voice"], "Cherry")
        self.assertTrue(url.endswith(".wav"))

    def test_missing_tts_model_raises(self):
        provider = _local_tts_provider(tts_model="")
        with self.assertRaises(NotImplementedError):
            provider.tts("文本")


class _FakeEmbedResponse:
    """假 HTTP 响应：供 embed 超时测试使用。"""

    def __init__(self, *, vectors):
        self.status_code = 200
        self._body = {"data": [{"embedding": vec} for vec in vectors]}

    def raise_for_status(self):
        pass

    def json(self):
        return self._body


class EmbedTimeoutTests(SimpleTestCase):
    """embed 超时：管理端配置优先，未配则默认 60s（本地 embed 冷启动慢，需可调大）。"""

    def _provider(self, timeout=None):
        from apps.ai.providers.openai_compat import OpenAICompatProvider

        return OpenAICompatProvider(
            {
                "api_key": "local-key",
                "base_url": "http://127.0.0.1:8000/v1",
                "embed_model": "bge-m3",
                "timeout": timeout,
            }
        )

    def test_embed_timeout_defaults_to_60(self):
        provider = self._provider()
        seen = {}

        def fake_post(url, headers=None, json=None, timeout=None):
            seen["timeout"] = timeout
            return _FakeEmbedResponse(vectors=[[0.1, 0.2]])

        with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
            vectors = provider.embed(["文本"])
        self.assertEqual(seen["timeout"], 60)
        self.assertEqual(vectors, [[0.1, 0.2]])

    def test_embed_timeout_prefers_provider_config(self):
        provider = self._provider(timeout=180)
        seen = {}

        def fake_post(url, headers=None, json=None, timeout=None):
            seen["timeout"] = timeout
            return _FakeEmbedResponse(vectors=[[0.1, 0.2]])

        with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
            provider.embed(["文本"])
        self.assertEqual(seen["timeout"], 180)


class IngestPptPagesTests(TestCase):
    """课件 PPT 逐页入向量库（ingest_ppt_pages）。

    课件是向量库的主要内容；失败必须静默（返回 0、不动旧数据、不改 PPT 状态），
    由 QA 侧原文拼接兜底；全链路禁用 mock 兜底。
    """

    def setUp(self):
        from apps.courses.models import Catalog, Course, PPTResource

        self.teacher = User.objects.create_user(
            username="ppt-ingest-teacher", password="StrongPass!123", role=User.Role.TEACHER
        )
        self.course = Course.objects.create(name="Java", teacher=self.teacher)
        self.catalog = Catalog.objects.create(
            course=self.course, title="第一章", is_published=True
        )
        self.ppt = PPTResource.objects.create(
            course=self.course,
            catalog=self.catalog,
            file_name="a.pptx",
            version=1,
            is_active=True,
            parse_status=PPTResource.ParseStatus.DONE,
            parsed_pages=[
                {"page": 1, "title": "第一章", "body": "Java 简介"},
                {"page": 2, "title": "", "body": ""},  # 空白页应跳过
                {"page": 3, "title": "第三章", "body": "面向对象"},
            ],
        )

    def _seed_old_chunks(self, count=2):
        KnowledgeChunk.objects.bulk_create(
            [
                KnowledgeChunk(
                    ppt=self.ppt,
                    course_id=self.course.id,
                    content=f"旧片段{i}",
                    page=i + 1,
                    embedding=[0.5] * 4,
                )
                for i in range(count)
            ]
        )

    def test_ingest_creates_chunk_per_page(self):
        provider = _FakeIngestEmbedProvider(dimension=4)
        self._seed_old_chunks()
        with patch("apps.ai.services.get_provider", return_value=provider):
            count = ingest_ppt_pages(self.ppt)

        self.assertEqual(count, 2)  # 空白页跳过
        chunks = list(self.ppt.chunks.order_by("page"))
        self.assertEqual([c.page for c in chunks], [1, 3])
        self.assertEqual(chunks[0].content, "第一章\nJava 简介")
        self.assertEqual(chunks[0].course_id, self.course.id)
        self.assertIsNone(chunks[0].material_id)
        for c in chunks:
            self.assertEqual(len(c.embedding), 4)
        # 探针语义：入库同样禁用 mock 兜底
        self.assertEqual(provider.allow_fallback_seen, [False])

    def test_ingest_without_parsed_pages_returns_zero(self):
        self.ppt.parsed_pages = []
        self.ppt.save(update_fields=["parsed_pages", "updated_at"])
        with patch("apps.ai.services.get_provider", return_value=_FakeIngestEmbedProvider()):
            self.assertEqual(ingest_ppt_pages(self.ppt), 0)
        self.assertEqual(self.ppt.chunks.count(), 0)

    def test_ingest_mock_provider_silently_skips(self):
        """未配置向量模型（回退模拟模式）：不入库、不报错、不动旧数据。"""
        from apps.courses.models import PPTResource

        self._seed_old_chunks()
        with patch("apps.ai.services.get_provider", return_value=MockProvider()):
            self.assertEqual(ingest_ppt_pages(self.ppt), 0)
        self.assertEqual(self.ppt.chunks.count(), 2)
        self.assertEqual(self.ppt.chunks.first().content, "旧片段0")
        self.ppt.refresh_from_db()
        self.assertEqual(self.ppt.parse_status, PPTResource.ParseStatus.DONE)

    def test_ingest_embed_error_keeps_old_chunks(self):
        provider = _FakeIngestEmbedProvider(raise_error=RuntimeError("503"))
        self._seed_old_chunks()
        with patch("apps.ai.services.get_provider", return_value=provider):
            self.assertEqual(ingest_ppt_pages(self.ppt), 0)
        self.assertEqual(self.ppt.chunks.count(), 2)
        self.assertEqual(self.ppt.chunks.first().content, "旧片段0")

    def test_ingest_count_mismatch_keeps_old_chunks(self):
        provider = _FakeIngestEmbedProvider(drop_last=True)
        self._seed_old_chunks()
        with patch("apps.ai.services.get_provider", return_value=provider):
            self.assertEqual(ingest_ppt_pages(self.ppt), 0)
        self.assertEqual(self.ppt.chunks.count(), 2)
        self.assertEqual(self.ppt.chunks.first().content, "旧片段0")


class _FakeIngestEmbedProvider:
    """ingest_ppt_pages 用的真实向量服务替身。"""

    name = "tongyi"
    embed_model = "text-embedding-v3"

    def __init__(self, dimension=4, raise_error=None, drop_last=False):
        self.dimension = dimension
        self.raise_error = raise_error
        self.drop_last = drop_last
        self.allow_fallback_seen = []

    def embed(self, texts, *, allow_fallback=True):
        self.allow_fallback_seen.append(allow_fallback)
        if self.raise_error is not None:
            raise self.raise_error
        vectors = [[0.25] * self.dimension for _ in texts]
        return vectors[:-1] if self.drop_last else vectors


class EmbedBatchTests(SimpleTestCase):
    """向量接口批量上限：超过 10 条分批请求、按顺序拼接；失败语义与不分批一致。"""

    def _provider(self):
        from apps.ai.providers.openai_compat import OpenAICompatProvider

        return OpenAICompatProvider(
            {
                "api_key": "local-key",
                "base_url": "http://127.0.0.1:8000/v1",
                "embed_model": "bge-m3",
            }
        )

    def test_batches_and_concatenates_in_order(self):
        provider = self._provider()
        calls = []

        def fake_post(url, headers=None, json=None, timeout=None):
            batch = json["input"]
            calls.append(list(batch))
            return _FakeEmbedResponse(vectors=[[float(len(t))] for t in batch])

        with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
            vectors = provider.embed([f"文本{i}" for i in range(23)])

        self.assertEqual([len(c) for c in calls], [10, 10, 3])
        self.assertEqual([v[0] for v in vectors], [float(len(f"文本{i}")) for i in range(23)])

    def test_failure_falls_back_to_mock_for_all_texts(self):
        """允许回退时任一批次失败 → 整批 mock（与不分批语义一致，不混杂真假向量）。"""
        provider = self._provider()

        def fake_post(url, headers=None, json=None, timeout=None):
            if len(json["input"]) > 1 and json["input"][0] == "文本10":
                raise requests_exception()
            return _FakeEmbedResponse(vectors=[[0.1] for _ in json["input"]])

        with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
            vectors = provider.embed([f"文本{i}" for i in range(12)])
        # 全部为 64 维 mock 向量
        self.assertEqual(len(vectors), 12)
        self.assertTrue(all(len(v) == 64 for v in vectors))

    def test_failure_raises_without_fallback(self):
        provider = self._provider()

        def fake_post(url, headers=None, json=None, timeout=None):
            raise requests_exception()

        with patch("apps.ai.providers.openai_compat.requests.post", side_effect=fake_post):
            with self.assertRaises(Exception):
                provider.embed([f"文本{i}" for i in range(12)], allow_fallback=False)

    def test_error_detail_includes_nested_error_message(self):
        """OpenAI 兼容网关的错误体嵌套在 error 对象里，也要解析出真实原因。"""
        from apps.ai.providers.openai_compat import _http_error_detail

        class _Resp:
            status_code = 400

            def json(self):
                return {"error": {"code": "InvalidParameter", "message": "batch size is invalid"}}

        detail = _http_error_detail(_Resp())
        self.assertIn("InvalidParameter", detail)
        self.assertIn("batch size is invalid", detail)


def requests_exception():
    import requests

    return requests.ConnectionError("boom")
