"""实验报告生成：DOCX（python-docx）+ PDF（reportlab）。

版式还原实验平台报告：实验名称/学生信息/班级/日期/步骤清单/评分明细/实验结论。
中文字体策略：
- 本地开发：优先系统字体（Windows 微软雅黑 / Linux Noto CJK / macOS PingFang）；
- 生产 Docker：fonts-noto-cjk（deploy/Dockerfile.backend 安装，fonts.conf 已映射）。
"""
import os
from datetime import datetime

from django.conf import settings
from django.utils import timezone

from .models import LabReport

# 常见中文候选字体（按平台择优），reportlab 需要可注册的 TTF/TTC
_FONT_CANDIDATES = [
    # Windows
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\simsun.ttc",
    # Linux（Docker 内 fonts-noto-cjk）
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
    # macOS
    "/System/Library/Fonts/PingFang.ttc",
]

# python-docx 用的字体名（Word 渲染时按名称查找查看方系统字体）
DOCX_FONT_CJK = "微软雅黑"
DOCX_FONT_EASTASIA = "微软雅黑"


def _find_font() -> str | None:
    for path in _FONT_CANDIDATES:
        if os.path.exists(path):
            return path
    return None


def _fmt_score(value) -> str:
    return "—" if value is None else str(value)


def build_report_context(submission) -> dict:
    """从 LabSubmission 聚合报告所需数据（学生/班级/步骤/评分/结论）。"""
    from apps.classroom.models import ClassStudent

    student = submission.student
    lab = submission.lab
    classroom = submission.schedule.classroom if submission.schedule_id else None
    # 学生在该课程下的第一个班级（报告展示用）
    if classroom is None:
        cs = ClassStudent.objects.filter(
            student=student, classroom__courses=lab.course, learn_status="active"
        ).select_related("classroom").first()
        classroom = cs.classroom if cs else None

    breakdown = submission.score_breakdown or {}
    steps = submission.steps_result or {}
    if isinstance(steps, list):  # 兼容列表式步骤
        steps = {str(i + 1): v for i, v in enumerate(steps)}
    step_lines = [
        f"步骤 {sid}：{'通过' if state == 'pass' else '未通过'}"
        for sid, state in sorted(steps.items(), key=lambda kv: (len(kv[0]), kv[0]))
    ]
    error_lines = [
        f"步骤 {e.get('step', '?')}：命令 {e.get('cmd', '')} {e.get('msg', '')}".strip()
        for e in (submission.error_log or [])
    ]
    return {
        "lab_title": lab.title,
        "course_name": lab.course.name,
        "student_name": student.real_name or student.username,
        "student_no": student.username,
        "class_name": classroom.name if classroom else "—",
        "experiment_date": submission.submitted_at or timezone.now(),
        "score": {
            "base": _fmt_score(submission.base_score),
            "accuracy": _fmt_score(submission.accuracy_score),
            "efficiency": _fmt_score(submission.efficiency_score),
            "completion": _fmt_score(submission.completion_score),
            "question": _fmt_score(submission.question_score),
            "total": _fmt_score(submission.total_score),
            "full": str(lab.total_score),
        },
        "step_lines": step_lines or ["（无步骤记录）"],
        "error_lines": error_lines,
        "elapsed_minutes": round(submission.elapsed_seconds / 60) if submission.elapsed_seconds else 0,
        "conclusion": "",
    }


def generate_report(submission) -> LabReport:
    """生成（或重新生成）实验报告，返回 LabReport 记录。幂等：同 submission 覆盖。"""
    from .models import LabReport

    report, _ = LabReport.objects.get_or_create(submission=submission)
    ctx = build_report_context(submission)
    ctx["conclusion"] = report.conclusion or "（学生未填写实验结论）"

    base_name = (
        f"实验报告_{ctx['student_no']}_{submission.lab.template.code}_"
        f"{ctx['experiment_date']:%Y%m%d}"
    )
    docx_path = _render_docx(ctx)
    pdf_path = _render_pdf(ctx)

    with open(docx_path, "rb") as f:
        report.file_docx.save(f"{base_name}.docx", f, save=False)
    with open(pdf_path, "rb") as f:
        report.file_pdf.save(f"{base_name}.pdf", f, save=False)
    report.file_name = f"{base_name}.docx"
    report.generated_at = timezone.now()
    report.save()
    # 清理临时文件
    for p in (docx_path, pdf_path):
        try:
            os.remove(p)
        except OSError:
            pass
    return report


def _render_docx(ctx: dict) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = DOCX_FONT_CJK
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_EASTASIA)

    title = doc.add_heading("", level=0)
    run = title.add_run("虚拟仿真实验报告")
    run.font.name = DOCX_FONT_CJK
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_EASTASIA)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(
        f"{ctx['course_name']} · {ctx['lab_title']}\n"
        f"姓名：{ctx['student_name']}    学号：{ctx['student_no']}    "
        f"班级：{ctx['class_name']}\n"
        f"实验日期：{ctx['experiment_date']:%Y年%m月%d日}"
    )

    def _section(label: str):
        h = doc.add_heading("", level=1)
        r = h.add_run(label)
        r.font.name = DOCX_FONT_CJK
        r.font.size = Pt(14)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_EASTASIA)

    _section("一、完成步骤")
    for line in ctx["step_lines"]:
        doc.add_paragraph(line, style="List Bullet")

    if ctx["error_lines"]:
        _section("二、错误尝试记录")
        for line in ctx["error_lines"]:
            doc.add_paragraph(line, style="List Number")

    _section("三、评分明细")
    s = ctx["score"]
    doc.add_paragraph(
        f"基础分：{s['base']} / 80\n"
        f"操作准确性：{s['accuracy']} / 5\n"
        f"完成效率：{s['efficiency']} / 5\n"
        f"完成度：{s['completion']} / 10\n"
        f"题目得分：{s['question']}\n"
        f"总分：{s['total']} / {s['full']}"
    )

    _section("四、实验结论")
    doc.add_paragraph(ctx["conclusion"])

    out = os.path.join(settings.MEDIA_ROOT, "tmp_report_docx")
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, f"report_{datetime.now().timestamp()}.docx")
    doc.save(path)
    return path


def _render_pdf(ctx: dict) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, ListFlowable, ListItem,
    )

    font_name = "ReportCJK"
    font_path = _find_font()
    if font_path:
        pdfmetrics.registerFont(TTFont(font_name, font_path, subfontIndex=0))
    else:
        font_name = "Helvetica"  # 兜底（无中文字体环境，生产 Docker 必有 noto）

    styles = {
        "title": ParagraphStyle("t", fontName=font_name, fontSize=20, leading=28, alignment=1, spaceAfter=6),
        "info": ParagraphStyle("i", fontName=font_name, fontSize=11, leading=18, alignment=1, spaceAfter=10),
        "h": ParagraphStyle("h", fontName=font_name, fontSize=14, leading=20, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", fontName=font_name, fontSize=11, leading=17),
    }

    def _esc(text: str) -> str:
        return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    flow = [
        Paragraph("虚拟仿真实验报告", styles["title"]),
        Paragraph(
            f"{_esc(ctx['course_name'])} · {_esc(ctx['lab_title'])}<br/>"
            f"姓名：{_esc(ctx['student_name'])}　学号：{_esc(ctx['student_no'])}　"
            f"班级：{_esc(ctx['class_name'])}<br/>"
            f"实验日期：{ctx['experiment_date']:%Y年%m月%d日}",
            styles["info"],
        ),
        Paragraph("一、完成步骤", styles["h"]),
        ListFlowable(
            [ListItem(Paragraph(_esc(line), styles["body"])) for line in ctx["step_lines"]],
            bulletType="bullet",
        ),
    ]
    if ctx["error_lines"]:
        flow.append(Paragraph("二、错误尝试记录", styles["h"]))
        flow.append(ListFlowable(
            [ListItem(Paragraph(_esc(line), styles["body"])) for line in ctx["error_lines"]],
            bulletType="1",
        ))
    s = ctx["score"]
    flow += [
        Paragraph("三、评分明细", styles["h"]),
        Paragraph(
            f"基础分：{s['base']} / 80<br/>操作准确性：{s['accuracy']} / 5<br/>"
            f"完成效率：{s['efficiency']} / 5<br/>完成度：{s['completion']} / 10<br/>"
            f"题目得分：{s['question']}<br/>"
            f"<b>总分：{s['total']} / {s['full']}</b>",
            styles["body"],
        ),
        Paragraph("四、实验结论", styles["h"]),
        Paragraph(_esc(ctx["conclusion"]).replace("\n", "<br/>"), styles["body"]),
        Spacer(1, 6 * mm),
    ]

    out = os.path.join(settings.MEDIA_ROOT, "tmp_report_pdf")
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, f"report_{datetime.now().timestamp()}.pdf")
    SimpleDocTemplate(path, pagesize=A4, title="虚拟仿真实验报告").build(flow)
    return path
