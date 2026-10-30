"""三类批量导入（班级 / 教师 / 学生）管道测试。

对应方案 §4 P2 准出 10 项，在 service 层与 HTTP 层各覆盖一遍：
模板生成、逐行校验、自动建班（含未勾选不建）、空班不创建（X-33）、
编码大小写归一（X-32）、重复上传幂等、错误行整体回滚、失败行/密码表可下载。
"""
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.classroom.models import ClassRoom, ClassStudent
from apps.users.models import User

from .imports import (
    KIND_CLASS,
    KIND_STUDENT,
    KIND_TEACHER,
    ImportFileError,
    ImportValidationError,
    annotate_rows,
    build_template,
    commit,
    infer_enrollment_year,
    parse_excel,
)

TEMPLATE_URL = "/api/admin-panel/imports/template/"
PREVIEW_URL = "/api/admin-panel/imports/preview/"
COMMIT_URL = "/api/admin-panel/imports/commit/"


def make_xlsx(headers: list, rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(name: str, content: bytes) -> SimpleUploadedFile:
    return SimpleUploadedFile(
        name, content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


CLASS_HEADERS = ["班级编码 *", "班级名称 *", "入学年 *", "学制 *", "状态"]
STUDENT_HEADERS = ["学号", "姓名", "班级编码", "班级名称", "入学年"]
TEACHER_HEADERS = ["工号", "姓名"]


def class_rows(*rows):
    return make_xlsx(CLASS_HEADERS, list(rows))


def student_rows(*rows):
    return make_xlsx(STUDENT_HEADERS, list(rows))


def teacher_rows(*rows):
    return make_xlsx(TEACHER_HEADERS, list(rows))


def read_sheet(content: bytes) -> list[list]:
    wb = load_workbook(BytesIO(content))
    ws = wb.worksheets[0]
    data = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return data


class ImportTestBase(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="p_admin_imp", password="pw123456!", role=User.Role.ADMIN
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.admin)

    def preview(self, content: bytes, kind: str, **extra):
        payload = {"kind": kind, "file": upload("t.xlsx", content), **extra}
        return self.client.post(PREVIEW_URL, payload, format="multipart")

    def do_commit(self, kind: str, rows: list, **options):
        return self.client.post(
            COMMIT_URL, {"kind": kind, "rows": rows, "options": options}, format="json"
        )

    def parsed(self, content: bytes, kind: str) -> list[dict]:
        return parse_excel(BytesIO(content), kind)


# ---------------------------------------------------------------------------
# 模板
# ---------------------------------------------------------------------------
class TemplateTests(ImportTestBase):
    def test_template_headers_and_tips_sheet(self):
        for kind, first_header in ((KIND_CLASS, "班级编码 *"), (KIND_TEACHER, "工号 *"),
                                   (KIND_STUDENT, "学号 *")):
            wb = load_workbook(BytesIO(build_template(kind)))
            self.assertEqual(wb.worksheets[0].title, {"class": "班级导入", "teacher": "教师导入",
                                                      "student": "学生导入"}[kind])
            self.assertEqual(wb.worksheets[0].cell(row=1, column=1).value, first_header)
            tips = wb.worksheets[1]
            self.assertEqual(tips.title, "填写说明")
            self.assertGreater(tips.max_row, 3)
            wb.close()

    def test_template_roundtrip_parses_back(self):
        """模板表头带 *，填完回传仍能被识别并落到正确字段。"""
        content = build_template(KIND_CLASS)
        # 空模板（只有表头）应报「没有数据行」，而不是表头缺列
        with self.assertRaises(ImportFileError) as ctx:
            parse_excel(BytesIO(content), KIND_CLASS)
        self.assertIn("有效数据行", str(ctx.exception))

        wb = load_workbook(BytesIO(content))
        wb.worksheets[0].append(["CS-2400", "计算机2400班", 2026, 3, "筹备中"])
        buf = BytesIO()
        wb.save(buf)
        parsed = parse_excel(BytesIO(buf.getvalue()), KIND_CLASS)
        self.assertEqual(parsed, [{
            "row_no": 2, "code": "CS-2400", "name": "计算机2400班",
            "enrollment_year": 2026, "duration_years": 3, "status": "筹备中",
        }])

    def test_template_endpoint_requires_admin(self):
        teacher = User.objects.create_user(
            username="p_teacher_imp", password="pw123456!", role=User.Role.TEACHER
        )
        client = APIClient()
        client.force_authenticate(user=teacher)
        resp = client.get(TEMPLATE_URL, {"kind": KIND_CLASS})
        self.assertEqual(resp.status_code, 403)

        resp = self.client.get(TEMPLATE_URL, {"kind": KIND_CLASS})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content[:2] == b"PK")
        resp = self.client.get(TEMPLATE_URL, {"kind": "unknown"})
        self.assertEqual(resp.status_code, 400)


# ---------------------------------------------------------------------------
# 班级导入
# ---------------------------------------------------------------------------
class ClassImportTests(ImportTestBase):
    def _ten_valid(self):
        return [[f"CS-{2400 + i}", f"计算机{2400 + i}班", 2026, 3, ""] for i in range(10)]

    def test_ten_valid_two_duplicate_one_invalid(self):
        rows = self._ten_valid()
        rows.append(["CS-2400", "重复编码班", 2026, 3, ""])
        rows.append(["cs-2405", "大小写重复", 2026, 3, ""])
        rows.append(["CS-2499", "", 2026, 3, ""])  # 缺班级名称
        content = class_rows(*rows)

        parsed = self.parsed(content, KIND_CLASS)
        self.assertEqual(len(parsed), 13)
        # 表头带 * 也要能匹配回字段
        self.assertEqual(parsed[0]["code"], "CS-2400")

        result = commit(KIND_CLASS, parsed, {"skip_invalid": True})
        self.assertEqual(result["created_classes"].__len__(), 10)
        self.assertEqual(len(result["failures"]), 1)
        self.assertEqual(result["skipped"], 2)
        self.assertEqual(ClassRoom.objects.count(), 10)
        self.assertTrue(ClassRoom.objects.filter(code="CS-2400").exists())

        # 失败行可下载，含「错误原因」列且指向缺名称那行
        sheet = read_sheet(result["failure_file"])
        self.assertEqual(sheet[0][-1], "错误原因")
        self.assertEqual(len(sheet), 2)
        self.assertIn("名称", str(sheet[1][-1]))

    def test_invalid_row_rolls_back_when_skip_invalid_off(self):
        rows = self._ten_valid()
        rows.append(["CS-2499", "", 2026, 3, ""])
        parsed = self.parsed(class_rows(*rows), KIND_CLASS)
        with self.assertRaises(ImportValidationError) as ctx:
            commit(KIND_CLASS, parsed, {"skip_invalid": False})
        self.assertEqual(len(ctx.exception.rows), 11)
        self.assertEqual(ClassRoom.objects.count(), 0)  # 整体回滚

    def test_duplicate_code_does_not_overwrite_existing(self):
        ClassRoom.objects.create(code="CS-2400", name="老名字", enrollment_year=2024, duration_years=4)
        parsed = self.parsed(class_rows(["CS-2400", "新名字", 2026, 3, ""]), KIND_CLASS)
        result = commit(KIND_CLASS, parsed, {"skip_invalid": False})
        self.assertEqual(result["created_classes"], [])
        self.assertEqual(result["skipped"], 1)
        room = ClassRoom.objects.get(code="CS-2400")
        self.assertEqual(room.name, "老名字")
        self.assertEqual(room.duration_years, 4)

    def test_preview_marks_status_without_writing(self):
        content = class_rows(["CS-2400", "计算机2400班", 2026, 3, ""],
                             ["CS-2400", "重复", 2026, 3, ""],
                             ["CS-2401", "", 2026, 3, ""])
        resp = self.preview(content, KIND_CLASS)
        self.assertEqual(resp.status_code, 200, resp.data)
        data = resp.json()["data"]
        self.assertEqual(data["summary"]["total"], 3)
        self.assertEqual(data["summary"]["ok"], 1)
        self.assertEqual(data["summary"]["skipped"], 1)
        self.assertEqual(data["summary"]["error"], 1)
        self.assertEqual(ClassRoom.objects.count(), 0)


# ---------------------------------------------------------------------------
# 教师导入
# ---------------------------------------------------------------------------
class TeacherImportTests(ImportTestBase):
    def test_import_and_reimport_is_idempotent(self):
        content = teacher_rows(["T2024001", "张老师"], ["T2024002", "李老师"])
        parsed = self.parsed(content, KIND_TEACHER)
        result = commit(KIND_TEACHER, parsed, {"skip_invalid": False})
        self.assertEqual(result["created_users"], 2)
        self.assertEqual(User.objects.filter(role=User.Role.TEACHER).count(), 2)

        teacher = User.objects.get(username="T2024001")
        self.assertTrue(teacher.must_change_password)
        self.assertTrue(teacher.check_password("Zk024001"))

        password_sheet = read_sheet(result["password_file"])
        self.assertEqual(password_sheet[0], ["工号", "姓名", "初始密码"])
        self.assertIn("Zk024001", [row[2] for row in password_sheet[1:]])

        # 再传一遍：全部 skipped，不新建也不报错
        again = commit(KIND_TEACHER, self.parsed(content, KIND_TEACHER), {"skip_invalid": False})
        self.assertEqual(again["created_users"], 0)
        self.assertEqual(again["skipped"], 2)
        self.assertEqual(User.objects.filter(role=User.Role.TEACHER).count(), 2)

    def test_short_username_and_missing_name_are_errors(self):
        parsed = self.parsed(teacher_rows(["T001", "短工号"], ["T2024002", ""]), KIND_TEACHER)
        result = annotate_rows(KIND_TEACHER, parsed)
        self.assertEqual(result["summary"]["error"], 2)
        self.assertIn("6", result["rows"][0]["message"])


# ---------------------------------------------------------------------------
# 学生导入：入班 / 自动建班
# ---------------------------------------------------------------------------
class StudentImportTests(ImportTestBase):
    def setUp(self):
        super().setUp()
        self.room = ClassRoom.objects.create(
            code="CS-2401", name="计算机2401班", enrollment_year=2024, duration_years=4
        )

    def test_into_existing_class(self):
        content = student_rows(["20240101", "王小明", "cs-2401", "", ""],
                               ["20240102", "李小红", "", "", ""])
        parsed = self.parsed(content, KIND_STUDENT)
        result = commit(KIND_STUDENT, parsed, {"skip_invalid": False})
        self.assertEqual(result["created_users"], 2)
        self.assertEqual(result["created_members"], 1)
        self.assertEqual(ClassStudent.objects.filter(classroom=self.room).count(), 1)
        self.assertEqual(ClassRoom.objects.count(), 1)  # 空编码那行不建班

    def test_auto_create_class_only_when_approved(self):
        content = student_rows(["20240201", "张三", "CS-2501", "计算机2501班", 2025])
        parsed = self.parsed(content, KIND_STUDENT)

        preview = annotate_rows(KIND_STUDENT, parsed, {})
        self.assertEqual(preview["summary"]["pending_class_count"], 1)
        self.assertEqual(preview["pending_classes"][0]["code"], "CS-2501")
        self.assertEqual(ClassRoom.objects.count(), 1)  # 预览不建班

        # 不勾选 → 提交阶段该行转 error
        with self.assertRaises(ImportValidationError):
            commit(KIND_STUDENT, parsed, {"skip_invalid": False})
        self.assertEqual(ClassRoom.objects.count(), 1)

        result = commit(KIND_STUDENT, parsed,
                        {"skip_invalid": False, "approved_classes": ["CS-2501"]})
        self.assertEqual(len(result["created_classes"]), 1)
        room = ClassRoom.objects.get(code="CS-2501")
        self.assertEqual(room.name, "计算机2501班")
        self.assertEqual(room.enrollment_year, 2025)
        self.assertEqual(result["created_members"], 1)

    def test_same_code_five_rows_creates_one_class(self):
        rows = [[f"2024030{i}", f"学生{i}", "CS-2601", "", ""] for i in range(5)]
        parsed = self.parsed(student_rows(*rows), KIND_STUDENT)
        result = commit(KIND_STUDENT, parsed,
                        {"skip_invalid": False, "approved_classes": ["CS-2601"]})
        self.assertEqual(len(result["created_classes"]), 1)
        self.assertEqual(result["created_members"], 5)
        self.assertEqual(ClassRoom.objects.filter(code="CS-2601").count(), 1)

    def test_code_with_all_rows_invalid_creates_no_class(self):
        """X-33：某编码下学生全错，不留下空壳班。"""
        rows = [["20240401", "", "CS-2701", "", ""],
                ["20240402", "", "CS-2701", "", ""]]
        parsed = self.parsed(student_rows(*rows), KIND_STUDENT)
        preview = annotate_rows(KIND_STUDENT, parsed, {})
        self.assertEqual(preview["pending_classes"], [])
        with self.assertRaises(ImportValidationError):
            commit(KIND_STUDENT, parsed,
                   {"skip_invalid": False, "approved_classes": ["CS-2701"]})
        self.assertFalse(ClassRoom.objects.filter(code="CS-2701").exists())

    def test_case_variants_do_not_create_two_classes(self):
        """X-32：CS-2801 与 cs-2801 视为同一个班。"""
        rows = [["20240501", "学生一", "CS-2801", "", ""],
                ["20240502", "学生二", "cs-2801", "", ""]]
        parsed = self.parsed(student_rows(*rows), KIND_STUDENT)
        preview = annotate_rows(KIND_STUDENT, parsed, {})
        self.assertEqual(len(preview["pending_classes"]), 1)

        result = commit(KIND_STUDENT, parsed,
                        {"skip_invalid": False, "approved_classes": ["cs-2801"]})
        self.assertEqual(len(result["created_classes"]), 1)
        self.assertEqual(ClassRoom.objects.filter(code="CS-2801").count(), 1)

    def test_default_enrollment_year_used_when_blank(self):
        rows = [["20240601", "学生甲", "CS-2901", "", ""]]
        parsed = self.parsed(student_rows(*rows), KIND_STUDENT)
        result = commit(KIND_STUDENT, parsed,
                        {"skip_invalid": False, "approved_classes": ["CS-2901"],
                         "default_enrollment_year": 2026, "default_duration_years": 3})
        self.assertEqual(len(result["created_classes"]), 1)
        room = ClassRoom.objects.get(code="CS-2901")
        self.assertEqual(room.enrollment_year, 2026)
        self.assertEqual(room.duration_years, 3)

    def test_name_falls_back_to_code(self):
        rows = [["20240602", "学生乙", "CS-2902", "", ""]]
        parsed = self.parsed(student_rows(*rows), KIND_STUDENT)
        commit(KIND_STUDENT, parsed, {"skip_invalid": False, "approved_classes": ["CS-2902"]})
        self.assertEqual(ClassRoom.objects.get(code="CS-2902").name, "CS-2902")

    def test_auto_created_class_enrollment_year_falls_back_when_blank(self):
        """行内与向导都不填入学年：自动建班按当前学期/自然年推导，不落 NULL。

        NULL 会让毕业预览把 graduate_year=None 视为「已到毕业年份」，
        一键整班毕业会误伤这批学生（生产现场问题）。
        """
        rows = [["20240603", "学生丙", "CS-2903", "", ""]]
        parsed = self.parsed(student_rows(*rows), KIND_STUDENT)
        commit(KIND_STUDENT, parsed, {"skip_invalid": False, "approved_classes": ["CS-2903"]})
        room = ClassRoom.objects.get(code="CS-2903")
        self.assertIsNotNone(room.enrollment_year)
        self.assertEqual(room.enrollment_year, infer_enrollment_year())
        self.assertIsNotNone(room.graduate_year)
        # 预览清单里的待建班也带推导值（管理员在预览页能看到的不再是空）
        preview = annotate_rows(
            KIND_STUDENT,
            self.parsed(student_rows(["20240604", "学生丁", "CS-2904", "", ""]), KIND_STUDENT),
            {},
        )
        self.assertEqual(preview["pending_classes"][0]["enrollment_year"], infer_enrollment_year())

    def test_second_upload_is_idempotent(self):
        content = student_rows(["20240701", "赵一", "CS-2401", "", ""],
                               ["20240702", "钱二", "CS-2401", "", ""])
        first = commit(KIND_STUDENT, self.parsed(content, KIND_STUDENT), {"skip_invalid": False})
        self.assertEqual(first["created_users"], 2)
        self.assertEqual(first["created_members"], 2)

        second = commit(KIND_STUDENT, self.parsed(content, KIND_STUDENT), {"skip_invalid": False})
        self.assertEqual(second["created_users"], 0)
        self.assertEqual(second["created_members"], 0)
        self.assertEqual(second["skipped"], 2)
        self.assertEqual(User.objects.filter(role=User.Role.STUDENT).count(), 2)
        self.assertEqual(ClassStudent.objects.count(), 2)

    def test_existing_student_without_class_is_reused(self):
        user = User.objects.create_user(
            username="20240801", password="pw123456!", role=User.Role.STUDENT, real_name="孙三"
        )
        parsed = self.parsed(student_rows(["20240801", "孙三", "CS-2401", "", ""]), KIND_STUDENT)
        result = commit(KIND_STUDENT, parsed, {"skip_invalid": False})
        self.assertEqual(result["created_users"], 0)
        self.assertEqual(result["created_members"], 1)
        self.assertEqual(User.objects.filter(role=User.Role.STUDENT).count(), 1)
        self.assertEqual(ClassStudent.objects.get(student=user).classroom_id, self.room.id)
        # 密码已不可知，不能瞎给：密码表写「沿用原密码」
        sheet = read_sheet(result["password_file"])
        self.assertEqual(sheet[1][3], "沿用原密码")

    def test_student_in_another_class_is_skipped(self):
        other = ClassRoom.objects.create(code="CS-2402", name="计算机2402班")
        user = User.objects.create_user(
            username="20240901", password="pw123456!", role=User.Role.STUDENT, real_name="周四"
        )
        ClassStudent.objects.create(classroom=other, student=user)
        parsed = self.parsed(student_rows(["20240901", "周四", "CS-2401", "", ""]), KIND_STUDENT)
        result = annotate_rows(KIND_STUDENT, parsed, {})
        self.assertEqual(result["rows"][0]["row_status"], "skipped")
        self.assertIn("计算机2402班", result["rows"][0]["message"])

        committed = commit(KIND_STUDENT, parsed, {"skip_invalid": False})
        self.assertEqual(committed["skipped"], 1)
        self.assertEqual(ClassStudent.objects.get(student=user).classroom_id, other.id)

    def test_username_of_other_role_is_rejected(self):
        User.objects.create_user(
            username="20241001", password="pw123456!", role=User.Role.TEACHER
        )
        parsed = self.parsed(student_rows(["20241001", "吴五", "", "", ""]), KIND_STUDENT)
        result = annotate_rows(KIND_STUDENT, parsed, {})
        self.assertEqual(result["rows"][0]["row_status"], "error")
        self.assertIn("其他角色", result["rows"][0]["message"])

    def test_class_in_graduated_status_blocks_assignment(self):
        room = ClassRoom.objects.create(code="CS-2000", name="已毕业班",
                                        status=ClassRoom.Status.GRADUATED)
        parsed = self.parsed(student_rows(["20241101", "郑六", "CS-2000", "", ""]), KIND_STUDENT)
        result = annotate_rows(KIND_STUDENT, parsed, {})
        self.assertEqual(result["rows"][0]["row_status"], "error")
        self.assertIn("已毕业", result["rows"][0]["message"])
        self.assertFalse(ClassStudent.objects.filter(classroom=room).exists())

    def test_auto_created_class_name_lands_in_password_sheet(self):
        """自动建班的学生，密码表「班级」列必须填新建班名（曾因 target_class_id 为空而空白）。"""
        parsed = self.parsed(student_rows(["20243001", "甲", "CS-3001", "计算机3001班", 2026],
                                         ["20243002", "乙", "CS-3001", "", ""]), KIND_STUDENT)
        result = commit(KIND_STUDENT, parsed,
                        {"skip_invalid": False, "approved_classes": ["CS-3001"]})
        sheet = read_sheet(result["password_file"])
        self.assertEqual([row[2] for row in sheet[1:]], ["计算机3001班", "计算机3001班"])

    def test_password_rule_and_class_name_in_sheet(self):
        parsed = self.parsed(student_rows(["20241234", "钱七", "CS-2401", "", ""]), KIND_STUDENT)
        result = commit(KIND_STUDENT, parsed, {"skip_invalid": False})
        sheet = read_sheet(result["password_file"])
        self.assertEqual(sheet[0], ["学号", "姓名", "班级", "初始密码"])
        self.assertEqual(sheet[1], ["20241234", "钱七", "计算机2401班", "Zk241234"])

    def test_bad_code_pattern_is_error(self):
        parsed = self.parsed(student_rows(["20241301", "王八", "CS_2401", "", ""]), KIND_STUDENT)
        result = annotate_rows(KIND_STUDENT, parsed, {})
        self.assertEqual(result["rows"][0]["row_status"], "error")
        self.assertIn("短横线", result["rows"][0]["message"])


# ---------------------------------------------------------------------------
# HTTP 全链（预览 → 提交）
# ---------------------------------------------------------------------------
class ImportHttpFlowTests(ImportTestBase):
    def test_missing_required_header_is_400(self):
        content = make_xlsx(["班级编码", "班级名称"], [["CS-2400", "计算机2400班"]])
        resp = self.preview(content, KIND_CLASS)
        self.assertEqual(resp.status_code, 400)
        self.assertIn("入学年", resp.json()["message"])

    def test_non_xlsx_rejected(self):
        resp = self.client.post(PREVIEW_URL, {
            "kind": KIND_CLASS,
            "file": SimpleUploadedFile("t.xls", b"not excel", content_type="application/vnd.ms-excel"),
        }, format="multipart")
        self.assertEqual(resp.status_code, 400)

    def test_class_preview_then_commit(self):
        content = class_rows(["CS-2400", "计算机2400班", 2026, 3, ""],
                             ["CS-2401", "计算机2401班", 2026, 3, "筹备中"])
        resp = self.preview(content, KIND_CLASS)
        self.assertEqual(resp.status_code, 200, resp.data)
        rows = resp.json()["data"]["rows"]
        self.assertEqual([r["row_status"] for r in rows], ["ok", "ok"])
        self.assertEqual(rows[1]["status_display"], "筹备中")

        resp = self.do_commit(KIND_CLASS, rows, skip_invalid=False)
        self.assertEqual(resp.status_code, 201, resp.data)
        data = resp.json()["data"]
        self.assertEqual(len(data["created_classes"]), 2)
        self.assertIsNone(data["failure_file"])
        self.assertIsNone(data["password_file"])
        self.assertEqual(ClassRoom.objects.get(code="CS-2401").status,
                         ClassRoom.Status.PREPARING)

    def test_student_preview_then_commit_with_approval(self):
        content = student_rows(["20242001", "甲", "CS-3001", "计算机3001班", 2026],
                               ["20242002", "乙", "CS-3001", "", ""])
        resp = self.preview(content, KIND_STUDENT)
        self.assertEqual(resp.status_code, 200, resp.data)
        data = resp.json()["data"]
        self.assertEqual(data["summary"]["pending_class_count"], 1)
        self.assertEqual(data["summary"]["pending_student_count"], 2)
        pending = data["pending_classes"][0]
        self.assertFalse(pending["approved"])
        rows = data["rows"]

        # 未勾选直接提交 → 400 且带逐行结果
        resp = self.do_commit(KIND_STUDENT, rows, skip_invalid=False, approved_classes=[])
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(ClassRoom.objects.count(), 0)

        resp = self.do_commit(KIND_STUDENT, rows, skip_invalid=False,
                              approved_classes=[pending["code"]])
        self.assertEqual(resp.status_code, 201, resp.data)
        payload = resp.json()["data"]
        self.assertEqual(payload["created_users"], 2)
        self.assertEqual(payload["created_members"], 2)
        self.assertEqual(len(payload["created_classes"]), 1)
        self.assertIsNotNone(payload["password_file"])
        self.assertEqual(payload["password_file"]["filename"], "initial_passwords_student.xlsx")

        decoded = read_sheet(__import__("base64").b64decode(
            payload["password_file"]["content_base64"]))
        self.assertEqual(decoded[0], ["学号", "姓名", "班级", "初始密码"])

    def test_commit_returns_failure_file(self):
        content = class_rows(["CS-2400", "计算机2400班", 2026, 3, ""],
                             ["CS-2401", "", 2026, 3, ""])
        rows = self.preview(content, KIND_CLASS).json()["data"]["rows"]
        resp = self.do_commit(KIND_CLASS, rows, skip_invalid=True)
        self.assertEqual(resp.status_code, 201, resp.data)
        data = resp.json()["data"]
        self.assertEqual(data["failed"], 1)
        self.assertIsNotNone(data["failure_file"])
        self.assertEqual(read_sheet(__import__("base64").b64decode(
            data["failure_file"]["content_base64"]))[0][-1], "错误原因")

    def test_teacher_cannot_use_import_endpoints(self):
        teacher = User.objects.create_user(
            username="p_teacher_imp2", password="pw123456!", role=User.Role.TEACHER
        )
        client = APIClient()
        client.force_authenticate(user=teacher)
        content = class_rows(["CS-2400", "计算机2400班", 2026, 3, ""])
        self.assertEqual(client.post(PREVIEW_URL, {
            "kind": KIND_CLASS, "file": upload("t.xlsx", content)}, format="multipart").status_code, 403)
        self.assertEqual(client.post(COMMIT_URL, {
            "kind": KIND_CLASS, "rows": [{"code": "CS-2400"}], "options": {}},
            format="json").status_code, 403)
