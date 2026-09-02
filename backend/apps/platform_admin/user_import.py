"""学生账号批量导入：模板生成、Excel 解析与行校验（预览与确认共用同一套校验）。"""
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db.models.functions import Lower
from openpyxl import Workbook, load_workbook

User = get_user_model()

# 初始密码规则：Lylg + 用户名后 6 位（不足 6 位取整个用户名）
PASSWORD_PREFIX = "Lylg"
# 用户名最少 4 位，保证初始密码 Lylg+用户名 至少 8 位（对齐系统密码长度约束）
MIN_USERNAME_LENGTH = 4
MAX_IMPORT_ROWS = 500
MAX_FILE_SIZE = 5 * 1024 * 1024

HEADER_USERNAME = "用户名"
HEADER_REAL_NAME = "姓名"


def initial_password_for(username: str) -> str:
    return f"{PASSWORD_PREFIX}{username[-6:]}"


def build_template() -> bytes:
    """生成导入模板：数据页仅表头，说明页写清填写规则与初始密码规则。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "学生导入"
    ws.append([HEADER_USERNAME, HEADER_REAL_NAME])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18

    tips = wb.create_sheet("填写说明")
    tips.append(["学生账号批量导入说明"])
    tips.append(["1. 请在「学生导入」页从第 2 行开始填写，每行一名学生，表头请勿修改。"])
    tips.append(["2. 用户名为学生登录账号（如学号），4 位以上，平台内不可重复。"])
    tips.append(["3. 初始密码由系统自动生成，规则为：Lylg + 用户名后 6 位（不足 6 位取整个用户名）。"])
    tips.append(["4. 例如用户名 2024012345 的初始密码为 Lylg012345。"])
    tips.append(["5. 学生首次登录后，系统会提示修改密码并补充手机号。"])
    tips.column_dimensions["A"].width = 90

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_text(row, index) -> str:
    """单元格统一转字符串：数字学号按整数处理，避免 2024010101.0 这类尾巴。"""
    if index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_excel(file_obj) -> list[dict]:
    """解析上传的 xlsx，返回 [{row_no, username, real_name}]；文件级错误抛 ValueError。"""
    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:
        raise ValueError("文件无法解析，请上传 .xlsx 格式的模板文件")
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise ValueError("Excel 内容为空")

    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    if HEADER_USERNAME not in header or HEADER_REAL_NAME not in header:
        raise ValueError("表头需包含「用户名」「姓名」两列，请下载系统模板填写")
    col_username = header.index(HEADER_USERNAME)
    col_real_name = header.index(HEADER_REAL_NAME)

    parsed = []
    for row_no, row in enumerate(rows[1:], start=2):
        username = _cell_text(row, col_username)
        real_name = _cell_text(row, col_real_name)
        if not username and not real_name:
            continue  # 跳过空行
        parsed.append({"row_no": row_no, "username": username, "real_name": real_name})

    if not parsed:
        raise ValueError("未读取到有效数据行，请从第 2 行开始填写")
    if len(parsed) > MAX_IMPORT_ROWS:
        raise ValueError(f"单次最多导入 {MAX_IMPORT_ROWS} 行，请拆分文件")
    return parsed


def annotate_rows(rows: list[dict]) -> list[dict]:
    """对每行做完整校验并附加 initial_password / status(ok|error) / message。

    预览与确认导入共用本函数，确认时以服务端校验结果为准（不信任前端编辑结果）。
    """
    cleaned = []
    for index, raw in enumerate(rows, start=1):
        username = raw.get("username")
        real_name = raw.get("real_name")
        cleaned.append({
            "row_no": raw.get("row_no") or index,
            "username": str(username).strip() if username is not None else "",
            "real_name": str(real_name).strip() if real_name is not None else "",
        })

    usernames = [item["username"] for item in cleaned if item["username"]]
    lowered = [name.lower() for name in usernames]
    # 用户名唯一性在整个系统按不区分大小写处理，这里保持一致
    existing_usernames = set(
        User.objects.annotate(lower_username=Lower("username"))
        .filter(lower_username__in=lowered)
        .values_list("lower_username", flat=True)
    ) if lowered else set()
    # 用户名同时不能与其他账号的手机号冲突（对齐后台新建账号的交叉校验）
    phone_clashes = set(
        User.objects.exclude(phone__isnull=True).exclude(phone="")
        .filter(phone__in=usernames)
        .values_list("phone", flat=True)
    ) if usernames else set()

    seen = {}
    duplicates = set()
    for item in cleaned:
        key = item["username"].lower()
        if not key:
            continue
        if key in seen:
            duplicates.add(key)
        else:
            seen[key] = True

    annotated = []
    for item in cleaned:
        username, real_name = item["username"], item["real_name"]
        message = ""
        if not username:
            message = "用户名不能为空"
        elif len(username) < MIN_USERNAME_LENGTH:
            message = f"用户名至少 {MIN_USERNAME_LENGTH} 位（用于生成初始密码）"
        elif len(username) > 150:
            message = "用户名过长"
        elif any(ch.isspace() for ch in username):
            message = "用户名不能包含空格"
        elif not real_name:
            message = "姓名不能为空"
        elif len(real_name) > 64:
            message = "姓名过长"
        elif username.lower() in duplicates:
            message = "与表内其他行用户名重复"
        elif username.lower() in existing_usernames:
            message = "该用户名已被使用"
        elif username in phone_clashes:
            message = "该用户名已被其他账号作为手机号使用"

        annotated.append({
            **item,
            "initial_password": initial_password_for(username) if username else "",
            "status": "error" if message else "ok",
            "message": message,
        })
    return annotated
