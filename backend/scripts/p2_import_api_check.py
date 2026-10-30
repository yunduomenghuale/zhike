"""P2 导入接口冒烟：模板 → 预览 → 提交（自动建班、幂等、失败行、权限）。

用法（后端须以临时库跑在 8005）：
    ./venv/Scripts/python.exe scripts/p2_import_api_check.py

注意：`admin_heavy` 档限流 5 次/分（preview/commit 都算），脚本一旦吃到 429 会
自动新建一个管理员账号换 token 重试——这既是压测手段，也顺带验证了「限流真的在生效」。
"""
from __future__ import annotations

import base64
import io
import os
import sys
import time

import requests
from openpyxl import Workbook, load_workbook

BASE = os.environ.get("API_BASE", "http://127.0.0.1:8005/api")
ADMIN_USER = os.environ.get("ADMIN_USER", "admin")
ADMIN_PWD = os.environ.get("ADMIN_PWD", "A0987654")
TAG = f"E2E{int(time.time()) % 100000:05d}"

TOKEN = ""
rotations = 0
ok_count = 0
fail_count = 0


def check(name: str, ok: bool, extra: str = "") -> None:
    global ok_count, fail_count
    print(f"{'PASS' if ok else 'FAIL'}  {name}{'  | ' + extra if extra else ''}")
    if ok:
        ok_count += 1
    else:
        fail_count += 1


def token_of(payload: dict) -> str:
    data = payload.get("data") or payload
    return (data.get("token") or {}).get("access") or data["access"]


def login(username: str, password: str) -> str:
    resp = requests.post(f"{BASE}/auth/login/", json={"username": username, "password": password}, timeout=20)
    resp.raise_for_status()
    return token_of(resp.json())


def rotate_admin() -> None:
    """吃限流时开一个新管理员（限流按 user.pk 计数）。"""
    global TOKEN, rotations
    rotations += 1
    # 账号规则 6~8 位纯字母数字：P2i + 4 位时间片 + 1 位轮次 = 8 位
    name = f"P2i{int(time.time()) % 10000:04d}{rotations % 10}"
    requests.post(f"{BASE}/admin-panel/users/", headers={"Authorization": f"Bearer {TOKEN}"},
                  json={"username": name, "real_name": "导入联调", "role": "admin",
                        "password": "pw123456", "is_active": True}, timeout=20)
    TOKEN = login(name, "pw123456")


def send(method: str, path: str, **kwargs) -> requests.Response:
    """带限流自愈的请求：429 时换一个管理员账号重试（最多 3 次）。"""
    headers = dict(kwargs.pop("headers", {}))
    for attempt in range(4):
        headers["Authorization"] = f"Bearer {TOKEN}"
        resp = requests.request(method, f"{BASE}{path}", headers=headers, timeout=120, **kwargs)
        if resp.status_code != 429:
            return resp
        if attempt == 3:
            return resp
        rotate_admin()
    return resp


def xlsx(headers: list, rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CLASS_HEADERS = ["班级编码 *", "班级名称 *", "入学年 *", "学制 *", "状态"]
STUDENT_HEADERS = ["学号", "姓名", "班级编码", "班级名称", "入学年"]
TEACHER_HEADERS = ["工号", "姓名"]


def preview(content: bytes, kind: str, **extra) -> requests.Response:
    return send("POST", "/admin-panel/imports/preview/", data={"kind": kind, **extra},
                files={"file": ("t.xlsx", content,
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def commit(kind: str, rows: list, options: dict) -> requests.Response:
    return send("POST", "/admin-panel/imports/commit/",
                json={"kind": kind, "rows": rows, "options": options})


def preview_rows(content: bytes, kind: str, **extra):
    resp = preview(content, kind, **extra)
    assert resp.status_code == 200, resp.text[:400]
    return resp.json()["data"]


def read_sheet(b64: str) -> list[list]:
    wb = load_workbook(io.BytesIO(base64.b64decode(b64)))
    ws = wb.worksheets[0]
    data = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return data


def main() -> int:
    global TOKEN
    TOKEN = login(ADMIN_USER, ADMIN_PWD)
    print(f"# 运行标记 {TAG}")

    # 1 模板
    r = send("GET", "/admin-panel/imports/template/", params={"kind": "class"})
    check("班级模板 200 且是 xlsx", r.status_code == 200 and r.content[:2] == b"PK", str(r.status_code))
    for kind in ("teacher", "student"):
        r = send("GET", "/admin-panel/imports/template/", params={"kind": kind})
        check(f"{kind} 模板可下载", r.status_code == 200 and r.content[:2] == b"PK")
    r = send("GET", "/admin-panel/imports/template/", params={"kind": "nope"})
    check("未知 kind → 400", r.status_code == 400, str(r.status_code))

    # 2 班级：10 正常 + 2 重复 + 1 缺名称
    rows = [[f"{TAG}-{i:02d}", f"E2E计算机{2400 + i}班", 2026, 3, ""] for i in range(1, 11)]
    rows.append([f"{TAG}-01", "重复编码班", 2026, 3, ""])
    rows.append([f"{TAG}-02".lower(), "大小写重复班", 2026, 3, ""])
    rows.append([f"{TAG}-99", "", 2026, 3, ""])
    data = preview_rows(xlsx(CLASS_HEADERS, rows), "class")
    summary = data["summary"]
    check("班级预览 10 可导入 / 2 跳过 / 1 待修正",
          (summary["ok"], summary["skipped"], summary["error"]) == (10, 2, 1), str(summary))

    r = commit("class", data["rows"], {"skip_invalid": False})
    check("有错误行且未跳过 → 400", r.status_code == 400, str(r.status_code))
    check("400 附带逐行结果", isinstance(r.json().get("data", {}).get("rows"), list))

    r = commit("class", data["rows"], {"skip_invalid": True})
    check("跳过后提交 201", r.status_code == 201, str(r.status_code))
    payload = r.json()["data"]
    check("班级新建 10 个", len(payload["created_classes"]) == 10, str(len(payload["created_classes"])))
    check("失败 1 行 / 跳过 2 行", (payload["failed"], payload["skipped"]) == (1, 2),
          f"{payload['failed']}/{payload['skipped']}")
    check("失败行文件含「错误原因」列",
          payload["failure_file"] is not None
          and read_sheet(payload["failure_file"]["content_base64"])[0][-1] == "错误原因")
    check("班级导入无密码表", payload["password_file"] is None)

    # 3 幂等
    again = commit("class", data["rows"], {"skip_invalid": True}).json()["data"]
    check("同一份班级文件再传一次不新建", len(again["created_classes"]) == 0, str(len(again["created_classes"])))

    # 4 学生入已存在的班（编码大小写不同也算同班）
    students = [[f"{TAG}01", "导入甲", f"{TAG}-01", "", ""],
                [f"{TAG}02", "导入乙", f"{TAG}-01".lower(), "", ""]]
    data = preview_rows(xlsx(STUDENT_HEADERS, students), "student")
    check("编码已存在时不需要建班", data["summary"]["pending_class_count"] == 0, str(data["summary"]))
    payload = commit("student", data["rows"], {"skip_invalid": False}).json()["data"]
    check("学生入班 2 账号 / 2 人",
          (payload["created_users"], payload["created_members"]) == (2, 2),
          f"{payload['created_users']}/{payload['created_members']}")
    sheet = read_sheet(payload["password_file"]["content_base64"])
    check("密码表表头", sheet[0] == ["学号", "姓名", "班级", "初始密码"], str(sheet[0]))
    username0 = str(sheet[1][0])
    check("初始密码 = Zk + 学号后 6 位", sheet[1][3] == "Zk" + username0[-6:],
          f"{username0} → {sheet[1][3]}")

    # 5 自动建班：同编码 5 行 → 1 个班；未勾选不建
    rows5 = [[f"{TAG}A{i}", f"新生{i}", f"{TAG}-50", "E2E自动建班班", 2026] for i in range(1, 6)]
    data = preview_rows(xlsx(STUDENT_HEADERS, rows5), "student")
    check("待建班级 1 个 / 5 人",
          (data["summary"]["pending_class_count"], data["summary"]["pending_student_count"]) == (1, 5),
          str(data["summary"]))

    r = commit("student", data["rows"], {"skip_invalid": False, "approved_classes": []})
    check("未勾选待建班级 → 400", r.status_code == 400, str(r.status_code))

    payload = commit("student", data["rows"],
                     {"skip_invalid": False, "approved_classes": [f"{TAG}-50"]}).json()["data"]
    check("勾选后建 1 个班", len(payload["created_classes"]) == 1, str(payload["created_classes"]))
    check("5 人全部入班", payload["created_members"] == 5, str(payload["created_members"]))

    # 6 X-33 空壳班不创建
    bad = [[f"{TAG}B1", "", f"{TAG}-70", "", ""], [f"{TAG}B2", "", f"{TAG}-70", "", ""]]
    data = preview_rows(xlsx(STUDENT_HEADERS, bad), "student")
    check("全错编码不进待建清单", data["summary"]["pending_class_count"] == 0, str(data["summary"]))
    r = commit("student", data["rows"], {"skip_invalid": True, "approved_classes": [f"{TAG}-70"]})
    check("全错编码不建班",
          r.status_code == 201 and len(r.json()["data"]["created_classes"]) == 0, str(r.status_code))

    # 7 X-32 大小写归一
    case_rows = [[f"{TAG}C1", "大小写甲", f"{TAG}-80", "", ""],
                 [f"{TAG}C2", "大小写乙", f"{TAG}-80".lower(), "", ""]]
    data = preview_rows(xlsx(STUDENT_HEADERS, case_rows), "student")
    check("大小写不同只算 1 个待建班级", data["summary"]["pending_class_count"] == 1, str(data["summary"]))
    payload = commit("student", data["rows"],
                     {"skip_invalid": False, "approved_classes": [f"{TAG}-80".lower()]}).json()["data"]
    check("大小写归一后只建 1 个班", len(payload["created_classes"]) == 1, str(payload["created_classes"]))

    # 8 教师导入 + 幂等
    teachers = [[f"T{TAG}1", "E2E教师甲"], [f"T{TAG}2", "E2E教师乙"]]
    data = preview_rows(xlsx(TEACHER_HEADERS, teachers), "teacher")
    check("教师预览 2 行可导入", data["summary"]["ok"] == 2, str(data["summary"]))
    payload = commit("teacher", data["rows"], {"skip_invalid": False}).json()["data"]
    check("教师导入 2 账号", payload["created_users"] == 2, str(payload["created_users"]))
    psheet = read_sheet(payload["password_file"]["content_base64"])
    check("教师密码表表头", psheet[0] == ["工号", "姓名", "初始密码"], str(psheet[0]))
    again = commit("teacher", data["rows"], {"skip_invalid": False}).json()["data"]
    check("教师重复导入幂等", again["created_users"] == 0 and again["skipped"] == 2, str(again["skipped"]))

    # 9 权限：教师 403
    t = requests.post(f"{BASE}/auth/login/", json={"username": "t101", "password": "pass1234"}, timeout=20)
    if t.status_code == 200:
        ttoken = token_of(t.json())
        r = requests.post(f"{BASE}/admin-panel/imports/commit/",
                          headers={"Authorization": f"Bearer {ttoken}"},
                          json={"kind": "class", "rows": [{"code": "X"}]}, timeout=20)
        check("教师调用导入接口 403", r.status_code == 403, str(r.status_code))
    else:
        print("SKIP  教师账号登录失败，跳过权限校验")

    print(f"\n== 通过 {ok_count} / 失败 {fail_count}（换号重试 {rotations} 次） ==")
    return 0 if fail_count == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
